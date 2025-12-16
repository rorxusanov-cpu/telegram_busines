from openpyxl import Workbook
from openpyxl.styles import Font
from database.db import cursor
from datetime import datetime
import os


def generate_excel(
    user_ids: list[int],
    date_from: str,
    date_to: str,
    filename: str
) -> str:
    # ❗ Agar user_ids bo‘sh bo‘lsa
    if not user_ids:
        raise ValueError("user_ids bo‘sh bo‘lishi mumkin emas")

    wb = Workbook()
    ws = wb.active
    ws.title = "Hisobot"

    headers = [
        "Sana",
        "Kim",
        "Amal",
        "Summa",
        "Valyuta",
        "Manba",
        "Izoh"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    placeholders = ",".join("?" for _ in user_ids)

    cursor.execute(f"""
        SELECT
            t.created_at,
            u.full_name,
            t.type,
            t.amount,
            t.currency,
            t.source,
            t.comment
        FROM transactions t
        JOIN users u ON u.id = t.user_id
        WHERE t.user_id IN ({placeholders})
          AND date(t.created_at) BETWEEN ? AND ?
        ORDER BY t.created_at
    """, (*user_ids, date_from, date_to))

    rows = cursor.fetchall()

    for created, name, typ, amount, currency, source, comment in rows:
        # Sana formatlash
        try:
            created_fmt = datetime.fromisoformat(created).strftime("%d.%m.%Y %H:%M")
        except Exception:
            created_fmt = created

        ws.append([
            created_fmt,
            name,
            "Kirim" if typ == "income" else "Chiqim",
            amount,
            currency,
            "Karta" if source in ("karta", "card") else "Naqd",
            comment or ""
        ])

    # Fayl yo‘li
    path = os.path.join("/mnt/data", filename)
    wb.save(path)

    return path
